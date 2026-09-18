# -*- coding: utf-8 -*-
"""不依赖任何真实工程数据的自测。

    pip install -r requirements.txt
    python -m unittest discover -s tests -v
"""
import datetime
import os
import shutil
import sys
import unittest

import openpyxl

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, REPO_ROOT)

from c600d import calc, config, ledger, writeback  # noqa: E402

BASE = datetime.date(2025, 11, 11)
LAYOUT = config.build_layout({})

#: 测试用临时目录放在仓库内（避免依赖系统 TEMP 的权限）
TMP_ROOT = os.path.join(REPO_ROOT, '.tmp_tests')


def fresh_dir(name):
    path = os.path.join(TMP_ROOT, name)
    shutil.rmtree(path, ignore_errors=True)
    os.makedirs(path, exist_ok=True)
    return path


def cn(d):
    return '%d年%02d月%02d日' % (d.year, d.month, d.day)


# --------------------------------------------------------------- 日记录解析
class TestParsing(unittest.TestCase):

    def test_parse_date_variants(self):
        self.assertEqual(config.parse_date('2026年04月06日'), datetime.date(2026, 4, 6))
        self.assertEqual(config.parse_date('2026-4-6'), datetime.date(2026, 4, 6))
        self.assertEqual(config.parse_date('2026/04/06'), datetime.date(2026, 4, 6))
        self.assertEqual(config.parse_date(datetime.datetime(2026, 4, 6, 8, 30)),
                         datetime.date(2026, 4, 6))
        self.assertEqual(config.parse_date(datetime.date(2026, 4, 6)),
                         datetime.date(2026, 4, 6))
        self.assertIsNone(config.parse_date('日期'))
        self.assertIsNone(config.parse_date(''))
        self.assertIsNone(config.parse_date(None))
        self.assertIsNone(config.parse_date('2026年13月40日'))   # 非法日期

    def test_parse_temp(self):
        self.assertEqual(config.parse_temp(10.5), 10.5)
        self.assertEqual(config.parse_temp(10), 10.0)
        self.assertEqual(config.parse_temp('9.4'), 9.4)          # 文本型数字
        self.assertEqual(config.parse_temp(-3.5), -3.5)
        self.assertIsNone(config.parse_temp(None))
        self.assertIsNone(config.parse_temp(''))
        self.assertIsNone(config.parse_temp('晴'))


# ------------------------------------------------------------------- 计算口径
class TestCalc(unittest.TestCase):

    def test_age_pouring_day_is_one(self):
        self.assertEqual(calc.compute_age(BASE, BASE), 1)
        self.assertEqual(calc.compute_age(BASE + datetime.timedelta(days=1), BASE), 2)
        # 与检验资料的等效龄期口径一致：制作日与检验日相差多少自然日，
        # 等效龄期即等于该天数（浇注当天计第 1 天，两者相减正好抵消）
        made = datetime.date(2025, 11, 12)
        tested = datetime.date(2026, 4, 8)
        self.assertEqual(calc.compute_age(tested, BASE) - calc.compute_age(made, BASE), 147)

    def test_cum_rule(self):
        self.assertEqual(calc.cum_rule(10.0), 10.0)
        self.assertEqual(calc.cum_rule(0.0), 0.0)      # T = 0 → 不累加
        self.assertEqual(calc.cum_rule(-3.5), 0.0)     # T < 0 → 按 0 计

    def _mk_cells(self, temps):
        cells = []
        for i, t in enumerate(temps):
            d = BASE + datetime.timedelta(days=i)
            cells.append(ledger.Cell('S1', 'left', 6 + i, d, t, None, None,
                                     'F', 'J', 'M'))
        return cells

    def test_cum_monotonic_and_frozen(self):
        # 5 天：10, -3, 0, 4, -1  →  10, 10, 10, 14, 14
        cells = self._mk_cells([10.0, -3.0, 0.0, 4.0, -1.0])
        results, _ = calc.compute_series(cells, BASE)
        self.assertEqual([r[2] for r in results], [10.0, 10.0, 10.0, 14.0, 14.0])
        cums = [r[2] for r in results]
        self.assertTrue(all(b >= a for a, b in zip(cums, cums[1:])))

    def test_age_series(self):
        cells = self._mk_cells([5.0] * 3)
        results, _ = calc.compute_series(cells, BASE)
        self.assertEqual([r[1] for r in results], [1, 2, 3])

    def test_cross_600_detection(self):
        # 100 天 × 6.0℃ = 600，应在第 100 天达标
        cells = self._mk_cells([6.0] * 100)
        _, cross = calc.compute_series(cells, BASE)
        self.assertEqual(cross['cum'], 600.0)
        self.assertEqual(cross['age'], 100)
        self.assertEqual(cross['date'], BASE + datetime.timedelta(days=99))

    def test_no_cross_returns_none(self):
        cells = self._mk_cells([1.0] * 10)
        _, cross = calc.compute_series(cells, BASE)
        self.assertIsNone(cross)

    def test_formula_same_sheet(self):
        cells = self._mk_cells([5.0, 6.0])
        self.assertEqual(calc.cum_formula(cells[0], cells[1]), '=J6+MAX(F7,0)')
        self.assertEqual(calc.age_formula(cells[0], cells[1]), '=M6+1')

    def test_formula_cross_sheet(self):
        a = ledger.Cell('S1', 'left', 27, BASE, 5.0, None, None, 'F', 'J', 'M')
        b = ledger.Cell('S2', 'left', 6, BASE + datetime.timedelta(days=1),
                        6.0, None, None, 'F', 'J', 'M')
        self.assertEqual(calc.cum_formula(a, b), '=S1!J27+MAX(F6,0)')
        self.assertEqual(calc.age_formula(a, b), '=S1!M27+1')

    def test_first_cell_formulas(self):
        c = ledger.Cell('S1', 'left', 6, BASE, 8.5, None, None, 'F', 'J', 'M')
        self.assertEqual(calc.first_value_formula(c, BASE), '=MAX(F6,0)')
        self.assertEqual(calc.first_age_formula(c, BASE), '=1')


# --------------------------------------------------------------- 端到端
class TestEndToEnd(unittest.TestCase):
    """用合成台账跑完整流程（不涉及任何真实工程数据）。"""

    N_DAYS = 50
    TEMPS = [10.0, -3.0, 0.0, 4.0, -1.0, 6.0, 12.5] * 8   # 56 项，取前 50

    def setUp(self):
        self.tmp = fresh_dir('e2e')
        self.path = os.path.join(self.tmp, 'ledger.xlsx')
        self._make_workbook()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _make_workbook(self):
        """造两页台账，每页左右各 22 行，共 44 天/页。"""
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        days = [BASE + datetime.timedelta(days=i) for i in range(self.N_DAYS)]
        for page in range(2):
            ws = wb.create_sheet('Sheet%d' % (page + 1))
            chunk = days[page * 44:(page + 1) * 44]
            for j, d in enumerate(chunk):
                side = 'left' if j < 22 else 'right'
                row = 6 + (j if side == 'left' else j - 22)
                cfg = LAYOUT[side]
                ws['%s%d' % (cfg['date'], row)] = cn(d)
                ws['%s%d' % (cfg['temp'], row)] = self.TEMPS[
                    (d - BASE).days % len(self.TEMPS)]
        wb.save(self.path)

    def test_read_validate_write_and_verify(self):
        cells = ledger.read_ledger(self.path, LAYOUT)
        self.assertEqual(len(cells), self.N_DAYS)
        self.assertEqual(ledger.validate(cells), [])

        results, cross = calc.compute_series(cells, BASE)
        # 独立重算一遍，两边必须一致
        run = 0.0
        for cell, age, cum in results:
            run = round(run + max(cell.temp, 0.0), 1)
            self.assertAlmostEqual(cum, run, places=6)
            self.assertEqual(age, (cell.date - BASE).days + 1)

        backup = os.path.join(self.tmp, 'backup.xlsx')
        stats = writeback.write_back(self.path, cells, LAYOUT, BASE,
                                     results, backup=backup)
        self.assertTrue(os.path.exists(backup))
        self.assertEqual(stats['cum_formulas'], self.N_DAYS - 1)
        self.assertEqual(stats['cached'], self.N_DAYS * 2)

        # 写入后：缓存值仍可读，且与算得一致
        wbv = openpyxl.load_workbook(self.path, data_only=True)
        wbf = openpyxl.load_workbook(self.path, data_only=False)
        for cell, age, cum in results:
            self.assertAlmostEqual(
                wbv[cell.sheet]['%s%d' % (cell.cum_col, cell.row)].value, cum, places=6)
            self.assertEqual(
                wbv[cell.sheet]['%s%d' % (cell.age_col, cell.row)].value, age)

        # 跨页那一格必须是跨表公式，而不是新的数值起点
        second = cells[44]
        f = wbf[second.sheet]['%s%d' % (second.cum_col, second.row)].value
        self.assertTrue(str(f).startswith('=Sheet1!'), '跨页公式异常：%r' % f)

        # 同一页内跨数据块也要连着，不重新起算
        edge = cells[22]
        f2 = wbf[edge.sheet]['%s%d' % (edge.cum_col, edge.row)].value
        self.assertTrue(str(f2).startswith('=J27+'), '跨块公式异常：%r' % f2)

    def test_temp_override(self):
        cells = ledger.read_ledger(self.path, LAYOUT)
        target = cells[3].date
        changed = ledger.apply_temp_overrides(cells, {target: 99.0})
        self.assertEqual(len(changed), 1)
        self.assertEqual(cells[3].temp, 99.0)

    def test_validate_reports_gap(self):
        wb = openpyxl.load_workbook(self.path)
        ws = wb['Sheet1']
        ws['B9'] = cn(BASE + datetime.timedelta(days=10))   # 人为制造断层
        wb.save(self.path)
        problems = ledger.validate(ledger.read_ledger(self.path, LAYOUT))
        self.assertTrue(any('不连续' in p for p in problems), problems)

    def test_validate_reports_missing_temp(self):
        wb = openpyxl.load_workbook(self.path)
        wb['Sheet1']['F6'] = None
        wb.save(self.path)
        problems = ledger.validate(ledger.read_ledger(self.path, LAYOUT))
        self.assertTrue(any('缺少平均温度' in p for p in problems), problems)


# --------------------------------------------------------------- 输入文件
class TestOverridesFile(unittest.TestCase):

    def setUp(self):
        self.tmp = fresh_dir('overrides')

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_csv_with_header_and_blanks(self):
        p = os.path.join(self.tmp, 't.csv')
        with open(p, 'w', encoding='utf-8') as fh:
            fh.write('日期,平均温度\n')
            fh.write('2025-11-11,8.5\n')
            fh.write('2025-11-12,\n')          # 空值 → 跳过
            fh.write('# 注释行\n')
            fh.write('2025-11-13,7.5\n')
        got = config.load_temp_overrides(p)
        self.assertEqual(got, {datetime.date(2025, 11, 11): 8.5,
                               datetime.date(2025, 11, 13): 7.5})

    def test_bad_date_raises(self):
        p = os.path.join(self.tmp, 'bad.csv')
        with open(p, 'w', encoding='utf-8') as fh:
            fh.write('2025-11-11,8.5\n')
            fh.write('不是日期,3.0\n')
        with self.assertRaises(config.ConfigError):
            config.load_temp_overrides(p)


if __name__ == '__main__':
    unittest.main(verbosity=2)
