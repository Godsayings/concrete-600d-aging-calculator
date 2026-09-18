# -*- coding: utf-8 -*-
"""台账读取：把工作簿解析成「按日期排序的日记录序列」。

每个工作表有左右两个数据块，块内每一行 = 一天。
本模块只负责读取与校验，不做任何写入。
"""
import datetime

import openpyxl

from .config import parse_date, parse_temp


class Cell:
    """一个日记录（对应台账里的一个数据行）。"""

    __slots__ = ('sheet', 'date', 'temp', 'cum', 'age', 'cum_col', 'age_col',
                 'temp_col', 'row', 'side')

    def __init__(self, sheet, side, row, date, temp, cum, age,
                 temp_col, cum_col, age_col):
        self.sheet = sheet
        self.side = side
        self.row = row
        self.date = date
        self.temp = temp
        self.cum = cum
        self.age = age
        self.temp_col = temp_col
        self.cum_col = cum_col
        self.age_col = age_col

    @property
    def cum_ref(self):
        return '%s!%s%d' % (self.sheet, self.cum_col, self.row)

    @property
    def temp_ref(self):
        return '%s!%s%d' % (self.sheet, self.temp_col, self.row)

    def __repr__(self):
        return '<Cell %s %s%d %s temp=%s>' % (
            self.sheet, self.cum_col, self.row, self.date, self.temp)


def _cell_value(ws, col, row):
    return ws['%s%d' % (col, row)].value


def read_ledger(path, layout, sheet_order=None):
    """读取台账，返回按日期升序排列的 Cell 列表。

    sheet_order: 显式指定工作表顺序；缺省用工作簿自身顺序。
    """
    wb = openpyxl.load_workbook(path, data_only=False)
    names = sheet_order or wb.sheetnames
    skip = set(layout.get('skip_sheets') or [])
    first = layout['first_data_row']
    nrows = layout['block_rows']

    cells = []
    for name in names:
        if name not in wb.sheetnames or name in skip:
            continue
        ws = wb[name]
        for side in ('left', 'right'):
            cfg = layout[side]
            for i in range(nrows):
                row = first + i
                raw_date = _cell_value(ws, cfg['date'], row)
                raw_temp = _cell_value(ws, cfg['temp'], row)
                d = parse_date(raw_date)
                if d is None and parse_temp(raw_temp) is None:
                    continue
                if d is None:
                    raise ValueError(
                        '%s %s第 %d 行有温度但没有可识别的日期：%r'
                        % (name, side, row, raw_date))
                cells.append(Cell(
                    sheet=name, side=side, row=row, date=d,
                    temp=parse_temp(raw_temp),
                    cum=_cell_value(ws, cfg['cum'], row),
                    age=_cell_value(ws, cfg['age'], row),
                    temp_col=cfg['temp'], cum_col=cfg['cum'], age_col=cfg['age'],
                ))
    wb.close()

    cells.sort(key=lambda c: (c.date, 0 if c.side == 'left' else 1))
    return cells


def validate(cells):
    """返回问题清单（空列表表示无问题）。"""
    problems = []
    if not cells:
        return ['未读取到任何日记录，请检查 layout 配置']
    seen = {}
    for c in cells:
        if c.date in seen:
            problems.append('日期重复：%s（%s 与 %s）' % (
                c.date, seen[c.date].cum_ref, c.cum_ref))
        seen[c.date] = c
    for a, b in zip(cells, cells[1:]):
        if (b.date - a.date).days != 1:
            problems.append('日期不连续：%s(%s) → %s(%s)，相隔 %d 天' % (
                a.date, a.cum_ref, b.date, b.cum_ref, (b.date - a.date).days))
    missing = [c.date for c in cells if c.temp is None]
    if missing:
        problems.append('以下日期缺少平均温度：%s' % ', '.join(str(d) for d in missing))
    return problems


def apply_temp_overrides(cells, overrides):
    """用外部温度（如气象网数据）替换指定日期的平均温度。

    返回被替换的 (日期, 原值, 新值) 列表。
    """
    if not overrides:
        return []
    by_date = {c.date: c for c in cells}
    changed = []
    for d, t in sorted(overrides.items()):
        c = by_date.get(d)
        if c is None:
            continue
        if c.temp is None or abs(c.temp - t) > 1e-9:
            changed.append((d, c.temp, t))
            c.temp = t
    return changed
